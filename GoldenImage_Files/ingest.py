#!/usr/bin/env python3
# =============================================================================
# ingest.py — Generic RAG Document Ingestion Pipeline
# Agent in a Box | Decision ID: d3a11098-1cf6-444f-8293-7462a6e9e3cc
#
# ARM-safe for Pi 5. No heavy ML libraries.
# Embeddings: nomic-embed-text via Ollama REST API (already on golden image).
# Vector store: ChromaDB (persistent, embedded mode).
# Fully idempotent — re-ingesting the same file is a no-op.
#
# Supported file types: .txt .md .pdf .docx .xlsx .csv
#
# Usage:
#   python ingest.py --docs /path/to/documents \
#                    --collection my_business  \
#                    --chroma-dir /path/to/chroma_db
# =============================================================================

import sys, hashlib, argparse
from pathlib import Path

try:
    import requests
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "requests", "-q",
                           "--break-system-packages"])
    import requests

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
OLLAMA_URL    = "http://localhost:11434"
EMBED_MODEL   = "nomic-embed-text"
CHUNK_SIZE    = 500   # characters per chunk
CHUNK_OVERLAP = 50    # overlap between chunks
SUPPORTED_EXT = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def ok(t):   print(f"  ✓  {t}")
def info(t): print(f"  →  {t}")
def warn(t): print(f"  ⚠  {t}", file=sys.stderr)


def chunk_text(text: str) -> list[str]:
    """Split text into overlapping fixed-size chunks."""
    chunks, start = [], 0
    text = text.strip()
    while start < len(text):
        chunks.append(text[start:start + CHUNK_SIZE])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


def get_embedding(text: str) -> list[float]:
    """Embed text using Ollama nomic-embed-text REST API."""
    r = requests.post(
        f"{OLLAMA_URL}/api/embeddings",
        json={"model": EMBED_MODEL, "prompt": text},
        timeout=60,
    )
    r.raise_for_status()
    return r.json()["embedding"]


# ---------------------------------------------------------------------------
# Text extraction (one function per format, graceful on missing deps)
# ---------------------------------------------------------------------------
def extract_txt(path: Path) -> str:
    return path.read_text(errors="ignore")


def extract_pdf(path: Path) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        warn(f"pypdf not installed — skipping {path.name}. Run: pip install pypdf")
        return ""
    except Exception as e:
        warn(f"PDF read error {path.name}: {e}")
        return ""


def extract_docx(path: Path) -> str:
    try:
        import docx
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    except ImportError:
        warn(f"python-docx not installed — skipping {path.name}. Run: pip install python-docx")
        return ""
    except Exception as e:
        warn(f"DOCX read error {path.name}: {e}")
        return ""


def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append(" | ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)
    except ImportError:
        warn(f"openpyxl not installed — skipping {path.name}. Run: pip install openpyxl")
        return ""
    except Exception as e:
        warn(f"XLSX read error {path.name}: {e}")
        return ""


EXTRACTORS = {
    ".txt":  extract_txt,
    ".md":   extract_txt,
    ".csv":  extract_txt,
    ".pdf":  extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".xls":  extract_xlsx,
}


def extract_text(path: Path) -> str:
    extractor = EXTRACTORS.get(path.suffix.lower())
    if not extractor:
        warn(f"Unsupported file type: {path.suffix} — skipping {path.name}")
        return ""
    return extractor(path)


# ---------------------------------------------------------------------------
# Core ingestion
# ---------------------------------------------------------------------------
def ingest_file(path: Path, collection, file_hash: str) -> int:
    """
    Ingest a single file. Returns number of chunks added.
    Idempotent: skips if file hash already exists in collection.
    """
    # Check if already indexed by hash
    existing = collection.get(where={"source_hash": {"$eq": file_hash}})
    if existing["ids"]:
        ok(f"Already indexed (skipping): {path.name}")
        return 0

    text = extract_text(path)
    if not text.strip():
        warn(f"No text extracted from {path.name} — skipping")
        return 0

    chunks = chunk_text(text)
    info(f"Indexing {path.name} → {len(chunks)} chunks…")

    ids, embeddings, metadatas, documents = [], [], [], []
    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        try:
            emb = get_embedding(chunk)
        except Exception as e:
            warn(f"Embedding error (chunk {i} of {path.name}): {e}")
            continue
        chunk_id = f"{file_hash}_{i}"
        ids.append(chunk_id)
        embeddings.append(emb)
        metadatas.append({
            "source":      path.name,
            "source_hash": file_hash,
            "chunk_index": i,
        })
        documents.append(chunk)

    if ids:
        collection.add(ids=ids, embeddings=embeddings,
                       metadatas=metadatas, documents=documents)
        ok(f"Indexed {len(ids)} chunks ← {path.name}")

    return len(ids)


def ingest_folder(docs_dir: Path, collection_name: str, chroma_dir: Path) -> int:
    """Ingest all supported files in docs_dir into ChromaDB. Returns total chunks."""
    try:
        import chromadb
    except ImportError:
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", "chromadb", "-q",
                               "--break-system-packages"])
        import chromadb

    client     = chromadb.PersistentClient(path=str(chroma_dir))
    collection = client.get_or_create_collection(collection_name)

    files = sorted(
        f for f in docs_dir.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXT
    )

    if not files:
        warn(f"No supported documents found in {docs_dir}")
        return 0

    print(f"\n  Found {len(files)} file(s) to process in {docs_dir}\n")
    total = 0
    for f in files:
        file_hash = hashlib.md5(f.read_bytes()).hexdigest()
        total += ingest_file(f, collection, file_hash)

    print(f"\n  Ingestion complete — {total} new chunks indexed into '{collection_name}'")
    return total


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="Ingest documents into ChromaDB using nomic-embed-text"
    )
    parser.add_argument("--docs",       required=True, help="Path to documents/ folder")
    parser.add_argument("--collection", required=True, help="ChromaDB collection name")
    parser.add_argument("--chroma-dir", required=True, help="ChromaDB persistence directory")
    args = parser.parse_args()

    docs_dir   = Path(args.docs)
    chroma_dir = Path(args.chroma_dir)

    if not docs_dir.exists():
        print(f"Documents folder not found: {docs_dir}")
        sys.exit(1)

    chroma_dir.mkdir(parents=True, exist_ok=True)
    ingest_folder(docs_dir, args.collection, chroma_dir)


if __name__ == "__main__":
    main()
